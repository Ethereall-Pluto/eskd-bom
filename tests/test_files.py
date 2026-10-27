import subprocess
import sys

import ezdxf
import pytest
from openpyxl import load_workbook

from eskd_bom import drawing, rules
from eskd_bom.models import Component, TitleBlock
from eskd_bom.reader import read_components
from eskd_bom.tables import read_csv, write_csv, write_table, write_xlsx

PARTS = [
    Component("R1", "Резистор 10 кОм"),
    Component("R2", "Резистор 10 кОм"),
    Component("C1", "Конденсатор 0,1 мкФ", "Не устанавливать"),
]


def make_schematic(path, parts, extra_blocks=0):
    doc = ezdxf.new("R2018")
    block = doc.blocks.new("EL")
    block.add_attdef("REF", (0, 0))
    block.add_attdef("NAME", (0, -3))
    block.add_attdef("NOTE", (0, -6))
    doc.blocks.new("FRAME")
    msp = doc.modelspace()
    for i, part in enumerate(parts):
        msp.add_blockref("EL", (i * 20, 0)).add_auto_attribs(
            {"REF": part.designator, "NAME": part.name, "NOTE": part.note}
        )
    for _ in range(extra_blocks):
        msp.add_blockref("FRAME", (0, 50))
    doc.saveas(path)


def texts(path):
    return [t.dxf.text for t in ezdxf.readfile(path).modelspace().query("TEXT")]


def test_reader_collects_attributes_and_skips_blocks_without_designator(tmp_path):
    path = tmp_path / "schematic.dxf"
    make_schematic(path, PARTS, extra_blocks=2)

    result = read_components(path)

    assert result.components == PARTS
    assert result.skipped_blocks == ["FRAME", "FRAME"]


def test_reader_fails_when_a_designated_block_has_no_name(tmp_path):
    path = tmp_path / "schematic.dxf"
    doc = ezdxf.new("R2018")
    doc.blocks.new("EL").add_attdef("REF", (0, 0))
    doc.modelspace().add_blockref("EL", (0, 0)).add_auto_attribs({"REF": "R1"})
    doc.saveas(path)

    with pytest.raises(ValueError, match="R1"):
        read_components(path)


def test_csv_round_trip(tmp_path):
    path = tmp_path / "parts.csv"
    write_csv(PARTS, path)
    assert read_csv(path) == PARTS


def test_xlsx_has_header_and_rows(tmp_path):
    path = tmp_path / "parts.xlsx"
    write_xlsx(PARTS, path)
    rows = list(load_workbook(path).active.values)
    assert rows[0] == ("Designator", "Name", "Note")
    assert rows[3] == ("C1", "Конденсатор 0,1 мкФ", "Не устанавливать")


def test_unknown_table_format_is_rejected(tmp_path):
    with pytest.raises(ValueError):
        write_table(PARTS, tmp_path / "parts.txt")


def test_render_draws_both_forms_and_numbers_the_sheets(tmp_path):
    components = [Component(f"R{i}", f"Резистор {i} кОм") for i in range(1, 41)]
    pages = rules.paginate(rules.build_rows(components))
    path = tmp_path / "pe3.dxf"

    drawing.render(pages, TitleBlock(code="АБВГ.468157.001ПЭ3", title="Плата"), path)
    found = texts(path)

    assert len(pages) == 2
    assert found.count("АБВГ.468157.001ПЭ3") == 2  # graph 2 on form 2 and on form 2a
    assert "Перечень элементов" in found
    assert "Разраб." in found and "Листов" in found
    assert {"R1", "R40", "Резисторы"} <= set(found)


def test_text_too_wide_for_its_cell_is_compressed(tmp_path):
    pages = rules.paginate(rules.build_rows([Component("VD1", "Диод"), Component("VD2", "Диод")]))
    path = tmp_path / "pe3.dxf"
    drawing.render(pages, TitleBlock(code="X", title="Очень длинное наименование изделия для графы один"), path)

    by_text = {t.dxf.text: t for t in ezdxf.readfile(path).modelspace().query("TEXT")}
    assert by_text["VD1, VD2"].dxf.width < 1
    assert by_text["Очень длинное наименование изделия для графы один"].dxf.width < 1


def test_cli_build_goes_from_schematic_to_parts_list(tmp_path):
    schematic, out = tmp_path / "schematic.dxf", tmp_path / "pe3.dxf"
    make_schematic(schematic, PARTS)

    result = subprocess.run(
        [sys.executable, "-m", "eskd_bom", "build", str(schematic), "-o", str(out), "--code", "X.001ПЭ3", "--title", "Плата"],
        capture_output=True, text=True,
    )

    assert result.returncode == 0, result.stderr
    assert "R1, R2" in texts(out)


def test_cli_reports_errors_without_a_traceback(tmp_path):
    bad = tmp_path / "bad.csv"
    bad.write_text("Designator,Name\nR1,Резистор\nR1,Резистор\n", encoding="utf-8")

    result = subprocess.run(
        [sys.executable, "-m", "eskd_bom", "generate", str(bad), "-o", str(tmp_path / "x.dxf"), "--code", "X", "--title", "Y"],
        capture_output=True, text=True,
    )

    assert result.returncode == 1
    assert result.stderr.startswith("error:") and "Traceback" not in result.stderr
